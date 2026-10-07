"""
pruebas/probar_logica.py
--------------------------
Prueba la LÓGICA del programa (sin interfaz): lotes (entradas, salidas de
un lote elegido, desechar caducados, limpiezas por lote, completar
servicios eligiendo lotes, guardar/cargar, sesiones antiguas) y
consumibles.

Es un script normal de Python (sin pytest). Se ejecuta en GitHub Actions
junto a las pruebas de la interfaz, y también a mano:
    python pruebas/probar_logica.py
"""

import io
import sys
import tempfile
from contextlib import redirect_stdout
from datetime import date, time, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "modulos"))

from inventario import Inventario, Producto  # noqa: E402
from metricas import Metricas  # noqa: E402
from recetario import Menu, Receta, Recetario  # noqa: E402
from servicios import Servicio  # noqa: E402
from exportador import exportar_todo  # noqa: E402
from servicios import RegistroServicios  # noqa: E402
from compras import GestorCompras  # noqa: E402

HOY = date.today()
fallos: list[str] = []

# Las pruebas de costes de servicios (hasta la sección de IVA) cuentan lo
# PAGADO, con IVA: el modo "no recupero el IVA". La sección de IVA prueba el otro.
from inventario import AJUSTES  # noqa: E402
AJUSTES["iva_recuperable"] = False


def comprobar(condicion: bool, descripcion: str) -> None:
    print(("✅ " if condicion else "❌ ") + descripcion)
    if not condicion:
        fallos.append(descripcion)


def silencio(funcion, *args, **kwargs):
    """Ejecuta sin llenar la pantalla con los print() de los módulos."""
    with redirect_stdout(io.StringIO()):
        return funcion(*args, **kwargs)


def inventario_secreto() -> Inventario:
    """Tu ejemplo: 1 kg de secreto que caduca pronto + una compra nueva que caduca más tarde."""
    inv = Inventario()
    silencio(inv.agregar_producto, Producto(
        "Secreto", "Carnes", 1, "kg", 14, "Carnicería Pepe", stock_minimo=0.5,
        fecha_caducidad=HOY + timedelta(days=10),
    ))
    silencio(inv.entrada_stock, "Secreto", 0.8, precio_unitario=15, proveedor="Ibéricos Sierra",
             fecha_caducidad=HOY + timedelta(days=15))
    return inv


print("--- Entradas: cada compra es un lote ---")
inv = inventario_secreto()
secreto = inv.buscar_producto("Secreto")
comprobar(len(secreto.lotes) == 2, "Una compra nueva crea un segundo lote")
comprobar(secreto.stock == 1.8, "El stock total es la suma de los lotes (1,8 kg)")
comprobar(secreto.fecha_caducidad == HOY + timedelta(days=10), "La caducidad del producto es la del lote más próximo")
comprobar(abs(secreto.valor_total() - 26) < 1e-9, "El valor es cada lote a su precio (14 + 12 = 26 €)")
comprobar(secreto.lotes[1].proveedor == "Ibéricos Sierra", "Cada lote guarda su proveedor")
silencio(inv.entrada_stock, "Secreto", 0.5, fecha_caducidad=HOY + timedelta(days=15))
comprobar(len(secreto.lotes) == 3, "Aunque tenga la misma caducidad, cada compra es un lote aparte")
comprobar(secreto.lotes[2].proveedor == "Carnicería Pepe" and secreto.lotes[2].precio_unitario == 15,
          "Sin indicar proveedor ni precio se usan el habitual y el de la última compra")
silencio(inv.agregar_producto, Producto("Aceite", "Aceites", 5, "litros", 4, "Oleícola"))
silencio(inv.entrada_stock, "Aceite", 5, precio_unitario=4.5)
aceite = inv.buscar_producto("Aceite")
comprobar(len(aceite.lotes) == 2 and aceite.fecha_caducidad is None,
          "Los productos sin caducidad también tienen un lote por compra")

print("\n--- Salidas: del lote que elige el usuario ---")
inv = inventario_secreto()
secreto = inv.buscar_producto("Secreto")
comprobar(silencio(inv.salida_stock, "Secreto", 0.3, "consumo", 2), "Se puede sacar del lote que caduca MÁS TARDE")
comprobar(abs(secreto.buscar_lote(2).cantidad - 0.5) < 1e-9 and secreto.buscar_lote(1).cantidad == 1,
          "Solo baja el lote elegido")
comprobar(not silencio(inv.salida_stock, "Secreto", 0.6, "consumo", 2), "No se puede sacar de un lote más de lo que tiene")
comprobar(inv.historial[-1].lote_id == 2 and inv.historial[-1].precio_unitario == 15,
          "El historial apunta el lote y su precio real")
silencio(inv.salida_stock, "Secreto", 0.5, "consumo", 2)
comprobar(secreto.buscar_lote(2) is None and len(secreto.lotes) == 1, "Un lote que llega a 0 desaparece")
silencio(inv.entrada_stock, "Secreto", 1)
comprobar(secreto.lotes[-1].id == 3, "Los números de lote no se reutilizan")

print("\n--- Desechar un lote caducado ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Tomate", "Verduras", 2, "kg", 2, "Huerta", fecha_caducidad=HOY - timedelta(days=1)))
silencio(inv.entrada_stock, "Tomate", 3, precio_unitario=2.5, fecha_caducidad=HOY + timedelta(days=5))
comprobar(len(inv.lotes_caducados()) == 1, "Se detecta el lote caducado")
comprobar(len(inv.lotes_proximos_a_caducar()) == 1, "Y el que caduca en los próximos días, aparte")
comprobar(silencio(inv.desechar_lote, "Tomate", 1), "Se desecha el lote caducado")
tomate = inv.buscar_producto("Tomate")
comprobar(tomate.stock == 3 and not inv.lotes_caducados(), "Solo queda el lote bueno")
desperdicio = Metricas(inv).cantidad_desperdiciada("Tomate", HOY, HOY)
comprobar(desperdicio == 2, "Queda registrado como desperdicio (2 kg)")
comprobar(Metricas(inv).gasto_por_categoria(HOY, HOY) == {"Verduras": 7.5},
          "El gasto cuenta la compra a su precio (3 kg x 2,5 €); el stock inicial no es una compra")

print("\n--- Limpieza de un lote concreto ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Pata", "Carnes", 2, "unidades", 45, "Carnicería Pepe",
                                        tiene_merma=True, peso_unitario=7))
silencio(inv.entrada_stock, "Pata", 1, precio_unitario=50, peso_unitario=8)
pata = inv.buscar_producto("Pata")
comprobar(pata.lotes[1].peso_unitario == 8, "Cada lote de patas guarda su propio peso por unidad")
try:
    silencio(inv.limpiar_producto, "Pata", 1, "Carne limpia", 5)
    comprobar(False, "Con varios lotes hay que elegir cuál se limpia")
except ValueError:
    comprobar(True, "Con varios lotes hay que elegir cuál se limpia")
limpieza = silencio(inv.limpiar_producto, "Pata", 1, "Carne limpia", 5, {"Hueso": 1.5}, lote_id=2)
comprobar(limpieza.peso_bruto_kg == 8 and limpieza.coste == 50, "Usa el peso y el precio de ESE lote (8 kg, 50 €)")
comprobar(limpieza.lote_origen and "Lote 2" in limpieza.lote_origen, "La limpieza recuerda de qué lote salió")
limpia = inv.buscar_producto("Carne limpia")
comprobar(limpia.lotes[0].precio_unitario == 10, "El lote limpio carga con todo el coste (50 € / 5 kg)")
comprobar(inv.buscar_producto("Hueso").lotes[0].precio_unitario == 0, "Los derivados entran a coste 0")
comprobar(pata.stock == 2 and pata.buscar_lote(2) is None, "Solo sale el lote limpiado")

print("\n--- Completar un servicio eligiendo lotes ---")
inv = inventario_secreto()
recetario = Recetario()
silencio(recetario.agregar_receta, Receta("Secreto a la brasa", "Principal", {"Secreto": 0.2}))
silencio(recetario.agregar_menu, Menu("Menú brasa", [recetario.recetas["Secreto a la brasa"]]))
servicio = Servicio(HOY, time(21, 0), 5, "Menú brasa")  # 1 kg de secreto

fila = recetario.previsualizar_consumo(servicio, inv)[0]
comprobar(fila["lotes_elegidos"] == [1] and fila["sin_asignar"] == 0,
          "Por defecto propone el lote que caduca antes (y aquí llega)")
fila = recetario.previsualizar_consumo(servicio, inv, {"Secreto": [2]})[0]
comprobar(abs(fila["sin_asignar"] - 0.2) < 1e-9 and fila["lotes_restantes"] == [1],
          "Si el lote elegido no llega, dice cuánto falta y con qué lotes se puede completar")
try:
    silencio(recetario.completar_servicio, servicio, inv, {"Secreto": [2]})
    comprobar(False, "No deja completar sin elegir con qué lote se completa")
except ValueError:
    comprobar(True, "No deja completar sin elegir con qué lote se completa")
comprobar(inv.buscar_producto("Secreto").stock == 1.8 and servicio.estado == "pendiente", "...y no toca nada")
silencio(recetario.completar_servicio, servicio, inv, {"Secreto": [2, 1]})
secreto = inv.buscar_producto("Secreto")
comprobar(secreto.buscar_lote(2) is None and abs(secreto.buscar_lote(1).cantidad - 0.8) < 1e-9,
          "Saca 0,8 del lote 2 y los 0,2 que faltan del lote 1")
comprobar(servicio.estado == "completado", "El servicio queda completado")

inv = inventario_secreto()
grande = Servicio(HOY, time(21, 0), 20, "Menú brasa")  # 4 kg: no hay suficiente
fila = recetario.previsualizar_consumo(grande, inv, {"Secreto": [2]})[0]
comprobar(fila["sin_asignar"] == 0 and fila["faltante"] == 2.2, "Si no hay bastante en total, se usan todos los lotes")
silencio(recetario.completar_servicio, grande, inv, {"Secreto": [2]})
comprobar(inv.buscar_producto("Secreto").stock == 0, "...y se descuenta todo lo que había")

print("\n--- Guardar, cargar y sesiones antiguas ---")
inv = inventario_secreto()
silencio(inv.salida_stock, "Secreto", 0.3, "consumo", 1)
copia = Inventario.from_dict(inv.to_dict())
s2 = copia.buscar_producto("Secreto")
comprobar([(l.id, l.cantidad, l.proveedor) for l in s2.lotes] == [(1, 0.7, "Carnicería Pepe"), (2, 0.8, "Ibéricos Sierra")],
          "Los lotes se guardan y se cargan igual")
comprobar(copia.historial[-1].lote == inv.historial[-1].lote, "El historial conserva el lote de cada movimiento")
antiguo = {"productos": [{
    "nombre": "Harina", "categoria": "Panadería", "stock": 3, "unidad": "kg", "precio_unitario": 1.2,
    "proveedor": "Harinas del Sur", "stock_minimo": 2, "fecha_caducidad": (HOY + timedelta(days=30)).isoformat(),
}], "historial": [{"producto_nombre": "Harina", "categoria": "Panadería", "tipo": "entrada", "cantidad": 3,
                   "unidad": "kg", "precio_unitario": 1.2, "fecha": HOY.isoformat()}]}
viejo = Inventario.from_dict(antiguo)
harina = viejo.buscar_producto("Harina")
comprobar(len(harina.lotes) == 1 and harina.stock == 3 and harina.precio_unitario == 1.2
          and harina.fecha_caducidad == HOY + timedelta(days=30),
          "Una sesión de antes de los lotes se convierte en un lote con su stock, precio y caducidad")

print("\n--- Corregir lotes y productos ---")
inv = inventario_secreto()
silencio(inv.editar_lote, "Secreto", 2, cantidad=0.75, fecha_caducidad=HOY + timedelta(days=20), proveedor="Otro")
lote = inv.buscar_producto("Secreto").buscar_lote(2)
comprobar(lote.cantidad == 0.75 and lote.proveedor == "Otro" and lote.fecha_caducidad == HOY + timedelta(days=20),
          "Se puede corregir la cantidad, el proveedor y la caducidad de un lote")
silencio(inv.editar_lote, "Secreto", 2, borrar_fecha_caducidad=True)
comprobar(lote.fecha_caducidad is None, "Se puede quitar la caducidad de un lote")
silencio(inv.editar_producto, "Secreto", nuevo_nombre="Secreto ibérico", stock_minimo=1)
comprobar(inv.buscar_producto("Secreto ibérico").stock == 1.75, "Renombrar el producto conserva sus lotes")

print("\n--- Excel ---")
inv = inventario_secreto()
with tempfile.TemporaryDirectory() as carpeta:
    ruta = silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta)
    from openpyxl import load_workbook
    libro = load_workbook(ruta)
    comprobar("Lotes" in libro.sheetnames and libro["Lotes"].max_row == 4,
              "El Excel tiene una hoja 'Lotes' con un renglón por lote (+ cabecera y total)")

print("\n--- Consumibles ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Pan", "Panadería", 5, "unidades", 0.5, "Horno"))
silencio(inv.agregar_producto, Producto(
    "Servilletas", "Desechables", 100, "unidades", 0.02, "Hostelería", stock_minimo=50,
    fecha_caducidad=HOY + timedelta(days=3), tipo="consumible",
))
serv_prod = inv.buscar_producto("Servilletas")
comprobar(serv_prod.es_consumible() and serv_prod.fecha_caducidad is None, "Un consumible no guarda caducidad")
try:
    Producto("Film", "Desechables", 1, "kg", 5, "Hostelería", tiene_merma=True, tipo="consumible")
    comprobar(False, "Un consumible no puede tener merma")
except ValueError:
    comprobar(True, "Un consumible no puede tener merma")
silencio(inv.entrada_stock, "Servilletas", 200, precio_unitario=0.01, fecha_caducidad=HOY + timedelta(days=3))
comprobar(serv_prod.lotes[-1].fecha_caducidad is None, "Sus compras tampoco guardan caducidad")
comprobar([p.nombre for p in inv.alimentos()] == ["Pan"] and [p.nombre for p in inv.consumibles()] == ["Servilletas"],
          "Alimentos y consumibles se listan por separado")

recetario = Recetario()
silencio(recetario.agregar_receta, Receta("Bocadillo", "Principal", {"Pan": 1}))
menu = Menu("Picnic", [recetario.recetas["Bocadillo"]], {"Servilletas": 2})
silencio(recetario.agregar_menu, menu)
comprobar(menu.ingredientes_por_comensal() == {"Pan": 1}, "A simple vista, el menú solo muestra sus ingredientes")
comprobar(menu.calcular_necesidades_totales(4) == {"Pan": 4, "Servilletas": 8}, "Las necesidades suman comida y consumibles")
comprobar(menu.costo_consumibles_por_comensal(inv) > 0 and menu.costo_por_comensal(inv) == 0.5,
          "El coste de la comida y el de los consumibles se calculan por separado")
servicio = Servicio(HOY, time(14, 0), 4, "Picnic")
silencio(recetario.completar_servicio, servicio, inv, {"Servilletas": [1]})
comprobar(serv_prod.stock == 292 and inv.buscar_producto("Pan").stock == 1,
          "Completar el servicio descuenta también los consumibles (8 servilletas)")
comprobar(inv.historial[-1].tipo_producto in ("alimento", "consumible")
          and any(m.tipo_producto == "consumible" and m.motivo == "consumo" for m in inv.historial),
          "El historial sabe qué movimientos son de consumibles")
gasto = Metricas(inv).gasto_por_tipo(HOY, HOY)
comprobar(gasto == {"alimento": 0, "consumible": 2.0, "mantenimiento": 0}, f"El gasto se separa en alimentos y consumibles ({gasto})")

gestor = GestorCompras()
grande = Servicio(HOY + timedelta(days=1), time(14, 0), 200, "Picnic")
silencio(gestor.generar_lista_desde_servicios, [grande], recetario, inv)
comprobar(any(i.ingrediente == "Servilletas" and i.cantidad == 108 for i in gestor.items),
          "La lista de la compra incluye los consumibles que faltan (400 - 292 = 108)")

actualizados = recetario.renombrar_producto("Servilletas", "Servilletas blancas")
comprobar("Servilletas blancas" in menu.consumibles_por_comensal and actualizados == ["menú Picnic"],
          "Renombrar un consumible lo actualiza en los menús")

copia = Recetario.from_dict(recetario.to_dict())
comprobar(copia.menus["Picnic"].consumibles_por_comensal == {"Servilletas blancas": 2},
          "Los consumibles del menú se guardan y se cargan")
comprobar(Inventario.from_dict(inv.to_dict()).buscar_producto("Servilletas").tipo == "consumible",
          "El tipo del producto se guarda y se carga")

silencio(inv.agregar_producto, Producto("Guantes", "Limpieza", 10, "unidades", 0.1, "Hostelería",
                                        fecha_caducidad=HOY + timedelta(days=90)))
silencio(inv.salida_stock, "Guantes", 2, "consumo", 1)
silencio(inv.editar_producto, "Guantes", tipo="consumible")
guantes = inv.buscar_producto("Guantes")
comprobar(guantes.es_consumible() and guantes.fecha_caducidad is None
          and all(m.tipo_producto == "consumible" for m in inv.historial if m.producto_nombre == "Guantes"),
          "Corregir el tipo a consumible quita la caducidad y pasa su historial a consumibles")

print("\n--- Gastos, precio de cobro y margen ---")
from gastos import Gasto, RegistroGastos, resumen_servicio  # noqa: E402
from persistencia import guardar_sesion, cargar_sesion  # noqa: E402
from metricas import ArchivoInformes  # noqa: E402

for mal, motivo in (
    (lambda: Gasto("Gasolina", "Transporte", 0), "importe 0"),
    (lambda: Gasto("Gasolina", "Gasolinera", 30), "categoría inventada"),
    (lambda: Gasto("  ", "Transporte", 30), "concepto vacío"),
):
    try:
        mal()
        comprobar(False, f"Un gasto con {motivo} se rechaza")
    except ValueError:
        comprobar(True, f"Un gasto con {motivo} se rechaza")

inv = inventario_secreto()
recetario = Recetario()
silencio(recetario.agregar_receta, Receta("Secreto a la brasa", "Principal", {"Secreto": 0.2}))
silencio(recetario.agregar_menu, Menu("Menú brasa", [recetario.recetas["Secreto a la brasa"]]))
registro = RegistroServicios()
con_precio = Servicio(HOY, time(21, 0), 4, "Menú brasa", precio_cobrado=100)
sin_precio = Servicio(HOY, time(14, 0), 2, "Menú brasa")
silencio(registro.agregar_servicio, con_precio)
silencio(registro.agregar_servicio, sin_precio)
gastos = RegistroGastos()
silencio(gastos.agregar_gasto, Gasto("Gasolina", "Transporte", 20, servicio_id=con_precio.id))
silencio(gastos.agregar_gasto, Gasto("Seguro furgoneta", "Seguros e impuestos", 60))

r = resumen_servicio(con_precio, inv, recetario, gastos)
comprobar(r["estimado"] and abs(r["comida"] - 0.8 * inv.buscar_producto("Secreto").precio_unitario) < 0.02
          and r["gastos"] == 20,
          "Servicio pendiente: coste estimado de la comida + sus gastos (no los generales)")
comprobar(r["margen"] == round(100 - r["coste_total"], 2), "Margen = cobro - coste")
r2 = resumen_servicio(sin_precio, inv, recetario, gastos)
comprobar(r2["cobrado"] is None and r2["margen"] is None, "Sin precio de cobro no hay margen (y no falla nada)")

silencio(recetario.completar_servicio, con_precio, inv, {"Secreto": [2]})
r = resumen_servicio(con_precio, inv, recetario, gastos)
comprobar(not r["estimado"] and abs(r["comida"] - (0.8 * 15)) < 1e-9,
          "Servicio completado: coste REAL de lo que salió (0,8 kg del lote a 15 €/kg = 12 €)")
comprobar(all(m.servicio_id == con_precio.id for m in inv.historial if m.motivo == "consumo"),
          "Las salidas de un servicio quedan apuntadas con su número")

with tempfile.TemporaryDirectory() as carpeta:
    ruta = str(Path(carpeta) / "sesion.json")
    silencio(guardar_sesion, inv, registro, recetario, GestorCompras(), ArchivoInformes(), ruta, gastos)
    sesion = silencio(cargar_sesion, ruta)
    comprobar(len(sesion.registro_gastos.gastos) == 2 and sesion.registro_servicios.servicios[0].precio_cobrado == 100
              and sesion.registro_servicios.servicios[1].precio_cobrado is None,
              "Los gastos y los precios de cobro se guardan y se cargan")
    comprobar(resumen_servicio(sesion.registro_servicios.servicios[0], sesion.inventario, sesion.recetario,
                               sesion.registro_gastos)["comida"] == r["comida"],
              "Tras guardar y cargar, el coste real del servicio sigue igual")
    import json
    with open(ruta, encoding="utf-8") as f:
        datos = json.load(f)
    del datos["gastos"]
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(datos, f)
    comprobar(silencio(cargar_sesion, ruta).registro_gastos.gastos == [],
              "Una sesión de antes de los gastos se carga sin gastos")
    ruta_excel = silencio(exportar_todo, inv, registro, GestorCompras(), carpeta, gastos, recetario)
    libro = load_workbook(ruta_excel)
    comprobar("Gastos" in libro.sheetnames and "Rentabilidad" in libro.sheetnames,
              "El Excel tiene las hojas 'Gastos' y 'Rentabilidad'")

print("\n--- Compra no prevista de un producto para un servicio ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Tomate", "Verduras", 0, "kg", 2, "Huerta"))
servicio = Servicio(HOY, time(14, 0), 10, "Menú")
servicio.id = 99
try:
    inv.compra_para_servicio("Tomate", 5, 6, 10, servicio.id)
    comprobar(False, "No se puede usar más de lo que se compró")
except ValueError:
    comprobar(inv.buscar_producto("Tomate").stock == 0, "No se puede usar más de lo que se compró (y no toca nada)")
lote = silencio(inv.compra_para_servicio, "Tomate", 5, 3, 12.5, servicio.id, "Mercado central",
                HOY + timedelta(days=4))
tomate = inv.buscar_producto("Tomate")
comprobar(tomate.stock == 2 and lote.precio_unitario == 2.5 and lote.proveedor == "Mercado central"
          and lote.fecha_caducidad == HOY + timedelta(days=4),
          "Entra como lote (a 12,5 € / 5 kg = 2,5 €/kg) y sobran 2 kg en el inventario")
servicio.estado = "completado"
r = resumen_servicio(servicio, inv, Recetario(), RegistroGastos())
comprobar(r["comida"] == 7.5, "El servicio carga solo lo usado (3 kg x 2,5 € = 7,5 €)")
comprobar(Metricas(inv).gasto_por_categoria(HOY, HOY) == {"Verduras": 12.5},
          "En Métricas cuenta como compra (12,5 €), una sola vez")

print("\n--- Lista de la compra al día ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Arroz", "Despensa", 1, "kg", 1.5, "Mayorista"))
silencio(inv.agregar_producto, Producto("Aceite", "Despensa", 0, "litros", 4, "Mayorista"))
recetario = Recetario()
silencio(recetario.agregar_receta, Receta("Paella", "Principal", {"Arroz": 0.1, "Aceite": 0.01}))
silencio(recetario.agregar_menu, Menu("Menú paella", [recetario.recetas["Paella"]]))
paella = Servicio(HOY + timedelta(days=2), time(14, 0), 50, "Menú paella")  # 5 kg de arroz, 0,5 l de aceite
gestor = GestorCompras()
silencio(gestor.generar_lista_desde_servicios, [paella], recetario, inv)
comprobar(gestor.pendiente_de("Arroz").cantidad == 4 and gestor.pendiente_de("Aceite") is not None,
          "La lista pide arroz (faltan 4 kg) y aceite")
silencio(inv.entrada_stock, "Arroz", 5, precio_unitario=1.4)  # comprado desde el inventario, sin pasar por la lista
avisos = silencio(gestor.generar_lista_desde_servicios, [paella], recetario, inv)
comprobar(gestor.pendiente_de("Arroz") is None and gestor.pendiente_de("Aceite") is not None,
          "Al regenerar la lista, se quita lo que ya no hace falta (el arroz) y se mantiene lo demás")
comprobar(any("Arroz" in a and "ya no hace falta" in a for a in avisos), "...y se avisa de ello")

print("\n--- Material reutilizable ---")
from materiales import Material, RegistroMaterial, lista_de_carga  # noqa: E402

try:
    Material("Plato", "Vajilla", 2.5)
    comprobar(False, "El material se cuenta en unidades enteras")
except ValueError:
    comprobar(True, "El material se cuenta en unidades enteras")

reg = RegistroMaterial()
silencio(reg.agregar_material, Material("Plato llano", "Vajilla", 50, 3.5))
silencio(reg.agregar_material, Material("Copa", "Cristalería", 30, 2.0))
comprobar(lista_de_carga({"Plato llano": 2, "Copa": 0.5}, 15) == {"Plato llano": 30, "Copa": 8},
          "La lista de carga se calcula por comensal y redondea hacia arriba (7,5 copas -> 8)")
silencio(reg.registrar_salida, 1, {"Plato llano": 30, "Copa": 8})
comprobar(reg.en_uso("Plato llano") == 30 and reg.disponibles("Plato llano") == 20
          and reg.buscar("Plato llano").cantidad_total == 50,
          "El material fuera sigue existiendo (50) pero no está disponible (20)")
try:
    reg.registrar_salida(2, {"Plato llano": 25})
    comprobar(False, "No se puede sacar más material del disponible")
except ValueError:
    comprobar(reg.en_uso("Plato llano") == 30, "No se puede sacar más material del disponible")
try:
    reg.editar_material("Plato llano", cantidad_total=10)
    comprobar(False, "El total no puede bajar de lo que está en uso")
except ValueError:
    comprobar(True, "El total no puede bajar de lo que está en uso")

incidencias = silencio(reg.registrar_vuelta, 1, {"Plato llano": 27, "Copa": 8}, {"Plato llano": 2})
comprobar(reg.en_uso("Plato llano") == 0 and reg.buscar("Plato llano").cantidad_total == 47,
          "Al volver: deja de estar en uso y lo que falta (3) deja de contar en el total")
comprobar(sorted((i.tipo, i.cantidad) for i in incidencias) == [("pérdida", 1), ("rotura", 2)],
          "De lo que falta: 2 roturas y 1 pérdida")
comprobar(reg.coste_incidencias_servicio(1) == 10.5, "Coste de roturas y pérdidas del servicio: 3 x 3,5 € = 10,5 €")

silencio(reg.reponer, "Plato llano", 3)
silencio(reg.dar_de_baja, "Copa", 1, "rotura")
comprobar(reg.buscar("Plato llano").cantidad_total == 50 and reg.buscar("Copa").cantidad_total == 29
          and reg.incidencias[-1].servicio_id is None,
          "Reponer suma unidades; una rotura en el almacén resta y no es de ningún servicio")

inv = inventario_secreto()
recetario = Recetario()
silencio(recetario.agregar_receta, Receta("Secreto a la brasa", "Principal", {"Secreto": 0.2}))
silencio(recetario.agregar_menu, Menu("Menú brasa", [recetario.recetas["Secreto a la brasa"]], {}, {"Plato llano": 2}))
servicio = Servicio(HOY, time(21, 0), 4, "Menú brasa", precio_cobrado=100)
servicio.id = 1
r = resumen_servicio(servicio, inv, recetario, RegistroGastos(), reg)
comprobar(r["material"] == 10.5 and r["coste_total"] == round(r["comida"] + 10.5, 2),
          "El coste del servicio incluye el material roto o perdido")

silencio(reg.editar_material, "Plato llano", nuevo_nombre="Plato blanco")
recetario.renombrar_material("Plato llano", "Plato blanco")
comprobar("Plato blanco" in reg.materiales and reg.salidas[0].cantidades.get("Plato blanco") == 30
          and reg.incidencias[0].material == "Plato blanco"
          and recetario.menus["Menú brasa"].materiales_por_comensal == {"Plato blanco": 2},
          "Renombrar un material lo actualiza en salidas, roturas y menús")

with tempfile.TemporaryDirectory() as carpeta:
    ruta = str(Path(carpeta) / "sesion.json")
    silencio(guardar_sesion, inv, RegistroServicios(), recetario, GestorCompras(), ArchivoInformes(), ruta,
             RegistroGastos(), reg)
    sesion = silencio(cargar_sesion, ruta)
    comprobar(sesion.registro_material.buscar("Plato blanco").cantidad_total == 50
              and len(sesion.registro_material.incidencias) == 3
              and sesion.recetario.menus["Menú brasa"].materiales_por_comensal == {"Plato blanco": 2},
              "El material, sus roturas y el material de los menús se guardan y se cargan")
    ruta_excel = silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta, RegistroGastos(),
                          recetario, reg)
    comprobar("Material" in load_workbook(ruta_excel).sheetnames, "El Excel tiene la hoja 'Material'")

print("\n--- Historial de servicios ---")
import historial  # noqa: E402

inv = inventario_secreto()
recetario = Recetario()
silencio(recetario.agregar_receta, Receta("Secreto a la brasa", "Principal", {"Secreto": 0.2}))
silencio(recetario.agregar_menu, Menu("Menú brasa", [recetario.recetas["Secreto a la brasa"]]))
registro = RegistroServicios()
garcia = Servicio(HOY, time(21, 0), 4, "Menú brasa", precio_cobrado=100, cliente="Familia García", lugar="Finca Los Olivos")
lopez = Servicio(HOY, time(14, 0), 15, "Menú brasa", cliente="López")  # pide 3 kg y solo hay 1,8
cancelado = Servicio(HOY, time(13, 0), 2, "Menú brasa", cliente="Ruiz")
for s_ in (garcia, lopez, cancelado):
    silencio(registro.agregar_servicio, s_)
cancelado.cancelar()
silencio(recetario.completar_servicio, garcia, inv, {"Secreto": [1]})
garcia.valoracion = "Sobró pan; llevar más hielo"
silencio(recetario.completar_servicio, lopez, inv)

comprobar(garcia.fecha_completado == HOY and garcia.menu_completado["recetas"][0]["nombre"] == "Secreto a la brasa",
          "Al completar se guardan la fecha y una copia del menú")
recetario.recetas["Secreto a la brasa"].ingredientes_por_comensal["Secreto"] = 0.5
comprobar(historial.previsto(garcia) == {"Secreto": 0.8},
          "Si después se cambia la receta, el historial sigue mostrando lo que se hizo (0,2 x 4 = 0,8 kg)")
filas = historial.previsto_frente_a_real(lopez, inv)
comprobar(filas[0]["previsto"] == 3 and abs(filas[0]["real"] - 1.0) < 1e-9 and abs(filas[0]["diferencia"] + 2) < 1e-9,
          "Previsto frente a real: el menú pedía 3 kg y solo salió 1 kg (lo que quedaba)")
comprobar([s_.id for s_ in historial.filtrar_servicios(registro, date(2000, 1, 1), date.max)] == [garcia.id, lopez.id],
          "El historial solo lista los completados, los más recientes primero")
comprobar(len(historial.filtrar_servicios(registro, date(2000, 1, 1), date.max, incluir_cancelados=True)) == 3,
          "Opcionalmente, también los cancelados")
comprobar([s_.id for s_ in historial.filtrar_servicios(registro, date(2000, 1, 1), date.max, cliente="Familia García")] == [garcia.id],
          "Filtrar por cliente")
comprobar([s_.id for s_ in historial.filtrar_servicios(registro, date(2000, 1, 1), date.max, texto="hielo")] == [garcia.id],
          "Buscar texto (también en la valoración)")
comprobar(historial.clientes(registro) == ["Familia García", "López", "Ruiz"], "Lista de clientes")
r = historial.resumen_periodo(historial.filtrar_servicios(registro, date(2000, 1, 1), date.max), inv, recetario,
                              RegistroGastos(), RegistroMaterial())
comprobar(r["servicios"] == 2 and r["comensales"] == 19 and r["facturado"] == 100 and r["servicios_sin_cobro"] == 1
          and r["menu_mas_repetido"] == "Menú brasa",
          "Resumen del periodo: 2 servicios, 19 comensales, 100 € facturados (1 sin precio)")
f = historial.ficha(garcia, inv, recetario, RegistroGastos(), RegistroMaterial())
comprobar(len(f["consumos"]) == 1 and f["consumos"][0]["lote"].startswith("Lote 1") and f["coste_por_comensal"] > 0,
          "La ficha recoge lo que salió de cada lote y el coste por comensal")
nuevo = silencio(historial.repetir_servicio, registro, garcia, HOY + timedelta(days=30), time(20, 0))
comprobar(nuevo.estado == "pendiente" and nuevo.menu == "Menú brasa" and nuevo.cliente == "Familia García"
          and nuevo.precio_cobrado == 100 and nuevo.comensales == 4 and nuevo.id != garcia.id,
          "Repetir un servicio crea uno nuevo igual en otra fecha")
copia = RegistroServicios.from_dict(registro.to_dict())
g2 = copia.buscar_por_id(garcia.id)
comprobar(g2.cliente == "Familia García" and g2.lugar == "Finca Los Olivos" and g2.valoracion == garcia.valoracion
          and g2.fecha_completado == HOY and g2.menu_completado == garcia.menu_completado,
          "Cliente, lugar, valoración y copia del menú se guardan y se cargan")
with tempfile.TemporaryDirectory() as carpeta:
    ruta_excel = silencio(exportar_todo, inv, registro, GestorCompras(), carpeta, RegistroGastos(), recetario, RegistroMaterial())
    hoja = load_workbook(ruta_excel)["Historial"]
    comprobar(hoja.max_row == 3 and hoja["D2"].value in ("Familia García", "López"),
              "El Excel tiene la hoja 'Historial' con una fila por servicio completado")

print("\n--- Elaboraciones (recetas preparadas por adelantado) ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Tomate", "Verduras", 2, "kg", 2, "Huerta", fecha_caducidad=HOY + timedelta(days=5)))
silencio(inv.entrada_stock, "Tomate", 3, precio_unitario=3, fecha_caducidad=HOY + timedelta(days=9))
silencio(inv.agregar_producto, Producto("Aceite", "Despensa", 5, "litros", 4, "Mayorista"))
recetario = Recetario()
ensalada = Receta("Ensalada", "Entrantes", {"Tomate": 0.1, "Aceite": 0.01}, vida_util_dias=3)
silencio(recetario.agregar_receta, ensalada)
silencio(recetario.agregar_menu, Menu("Menú ensalada", [ensalada]))
comprobar(ensalada.caducidad_propuesta(HOY) == HOY + timedelta(days=3), "La vida útil propone la caducidad (hoy + 3 días)")
comprobar(Recetario.from_dict(recetario.to_dict()).recetas["Ensalada"].vida_util_dias == 3, "La vida útil se guarda y se carga")

try:
    recetario.preparar_elaboracion("Ensalada", 100, inv)  # 10 kg de tomate: solo hay 5
    comprobar(False, "Sin ingredientes suficientes no se prepara")
except ValueError as e:
    comprobar("Tomate" in str(e) and inv.buscar_producto("Tomate").stock == 5 and not inv.elaboraciones.tandas,
              "Sin ingredientes suficientes no se prepara (y no toca nada)")
tanda = silencio(recetario.preparar_elaboracion, "Ensalada", 20, inv, {"Tomate": [2]}, ensalada.caducidad_propuesta(HOY))
tomate = inv.buscar_producto("Tomate")
comprobar(abs(tomate.buscar_lote(2).cantidad - 1) < 1e-9 and tomate.buscar_lote(1).cantidad == 2,
          "Preparar gasta los ingredientes del lote elegido (2 kg del lote 2)")
comprobar(inv.historial[-1].motivo == "elaboración", "Queda en el historial con el motivo 'elaboración'")
comprobar(abs(tanda.coste_por_racion - (2 * 3 + 0.2 * 4) / 20) < 1e-9 and tanda.raciones == 20
          and tanda.fecha_caducidad == HOY + timedelta(days=3),
          "La tanda tiene 20 raciones, su caducidad y su coste real por ración (0,34 €)")

servicio = Servicio(HOY + timedelta(days=1), time(14, 0), 30, "Menú ensalada")
servicio.id = 501
plan = recetario.plan_elaboraciones(servicio, inv)
comprobar(plan["Ensalada"]["raciones"] == 20 and plan["Ensalada"]["restantes"] == 10,
          "Por defecto se usan las raciones preparadas (20) y el resto (10) se hace con ingredientes")
filas = recetario.previsualizar_consumo(servicio, inv, {"Tomate": [1]})
comprobar(next(f for f in filas if f["ingrediente"] == "Tomate")["necesario"] == 1,
          "Solo se descuentan los ingredientes de las 10 raciones restantes (1 kg de tomate)")
sin_usar = recetario.previsualizar_consumo(servicio, inv, {"Tomate": [1]}, {"Ensalada": []})
comprobar(next(f for f in sin_usar if f["ingrediente"] == "Tomate")["necesario"] == 3,
          "Se puede elegir no usar las raciones preparadas (todo con ingredientes: 3 kg)")

gestor = GestorCompras()
avisos = silencio(gestor.generar_lista_desde_servicios,
                  [Servicio(HOY + timedelta(days=1), time(14, 0), 80, "Menú ensalada")], recetario, inv)
comprobar(gestor.pendiente_de("Tomate").cantidad == 3 and any("raciones preparadas" in a for a in avisos),
          "La lista de la compra descuenta lo ya preparado (80 - 20 = 60 raciones: 6 kg, hay 3 -> faltan 3)")
avisos = silencio(gestor.generar_lista_desde_servicios,
                  [Servicio(HOY + timedelta(days=10), time(14, 0), 80, "Menú ensalada")], recetario, inv)
comprobar(gestor.pendiente_de("Tomate").cantidad == 5,
          "...pero no cuenta una tanda que habrá caducado el día del servicio")

silencio(recetario.completar_servicio, servicio, inv, {"Tomate": [1]})
comprobar(not inv.elaboraciones.tandas and inv.elaboraciones.coste_servicio(501) == round(20 * tanda.coste_por_racion, 2),
          "Al completar, la tanda se gasta y su coste queda apuntado en el servicio")
r = resumen_servicio(servicio, inv, recetario, RegistroGastos())
comprobar(abs(r["comida"] - (20 * tanda.coste_por_racion + 1 * 2 + 0.1 * 4)) < 0.02,
          "El coste del servicio = raciones preparadas + ingredientes del resto")
f = historial.ficha(servicio, inv, recetario, RegistroGastos(), RegistroMaterial())
comprobar(any(c["producto"] == "Ensalada (preparada)" and c["cantidad"] == 20 for c in f["consumos"])
          and any(p["producto"] == "Tomate" and p["previsto"] == 1 and abs(p["diferencia"]) < 1e-9
                  for p in f["previsto_frente_a_real"]),
          "El historial muestra las raciones preparadas usadas y el previsto ya las descuenta")

vieja = silencio(recetario.preparar_elaboracion, "Ensalada", 5, inv, None, HOY - timedelta(days=1))
comprobar(inv.elaboraciones.caducadas() == [vieja], "Se detectan las tandas caducadas")
silencio(inv.elaboraciones.desechar, vieja.id)
comprobar(not inv.elaboraciones.tandas and Metricas(inv).valor_desperdiciado_total(HOY, HOY) == round(5 * vieja.coste_por_racion, 2),
          "Desechar una tanda cuenta como desperdicio, con su coste")
comprobar(Metricas(inv).cantidad_consumida("Tomate", HOY, HOY) > 0, "Lo gastado al preparar cuenta como consumo en Métricas")
silencio(recetario.preparar_elaboracion, "Ensalada", 4, inv, None, HOY + timedelta(days=2))
copia = Inventario.from_dict(inv.to_dict())
comprobar(len(copia.elaboraciones.tandas) == 1 and copia.elaboraciones.tandas[0].raciones == 4
          and len(copia.elaboraciones.usos) == len(inv.elaboraciones.usos),
          "Las tandas y sus usos se guardan y se cargan")

print("\n--- Elaboraciones base (sofritos, fondos, salsas...) ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Cebolla", "Verduras", 4, "kg", 1.5, "Huerta", fecha_caducidad=HOY + timedelta(days=10)))
silencio(inv.agregar_producto, Producto("Aceite", "Despensa", 5, "litros", 4, "Mayorista"))
silencio(inv.agregar_producto, Producto("Huesos", "Carnes", 3, "kg", 2, "Carnicería", fecha_caducidad=HOY + timedelta(days=4)))
silencio(inv.agregar_producto, Producto("Servilletas", "Mesa", 100, "unidades", 0.05, "Mayorista", tipo="consumible"))
sofrito = silencio(inv.definir_base, "Sofrito", "Elaboraciones", "kg", 1, {"Cebolla": 1.5, "Aceite": 0.1}, vida_util_dias=4)
comprobar(sofrito.es_base() and sofrito.stock == 0 and sofrito.tipo_descripcion() == "Elaboración base"
          and sofrito in inv.alimentos(), "Una elaboración base es un alimento más, con su fórmula y stock 0")
comprobar(sofrito.ingredientes_para(2) == {"Cebolla": 3, "Aceite": 0.2}, "La fórmula se escala a la cantidad que se prepara")
for malos, motivo in (
    ({"Patata": 1}, "un ingrediente que no existe"),
    ({"Servilletas": 1}, "un consumible"),
    ({"Sofrito": 1}, "a sí misma"),
):
    try:
        silencio(inv.editar_formula, "Sofrito", 1, malos)
        comprobar(False, f"La fórmula no admite {motivo}")
    except ValueError:
        comprobar(inv.buscar_producto("Sofrito").formula["ingredientes"] == {"Cebolla": 1.5, "Aceite": 0.1},
                  f"La fórmula no admite {motivo} (y no cambia nada)")
salsa = silencio(inv.definir_base, "Salsa", "Elaboraciones", "litros", 2, {"Sofrito": 0.5, "Huesos": 1})
try:
    silencio(inv.editar_formula, "Sofrito", 1, {"Cebolla": 1.5, "Salsa": 0.2})
    comprobar(False, "Las fórmulas no pueden ir en círculo")
except ValueError:
    comprobar(True, "Las fórmulas no pueden ir en círculo (Sofrito lleva Salsa, que lleva Sofrito)")

try:
    recetario = Recetario()
    recetario.preparar_base("Sofrito", 10, 9, inv)  # 15 kg de cebolla: solo hay 4
    comprobar(False, "Sin ingredientes suficientes no se prepara la base")
except ValueError as e:
    comprobar("Cebolla" in str(e) and sofrito.stock == 0, "Sin ingredientes suficientes no se prepara la base")
prep = silencio(recetario.preparar_base, "Sofrito", 2, 1.6, inv, {"Cebolla": [1], "Aceite": [1]})
comprobar(inv.buscar_producto("Cebolla").stock == 1 and abs(inv.buscar_producto("Aceite").stock - 4.8) < 1e-9,
          "Preparar 2 kg gasta lo que pide la fórmula (3 kg de cebolla, 0,2 l de aceite)")
comprobar(sofrito.stock == 1.6 and prep.prevista == 2 and prep.obtenida == 1.6 and prep.diferencia == -0.4,
          "Entra lo que salió DE VERDAD (1,6 kg) y queda apuntado: previsto 2, obtenido 1,6, diferencia -0,4")
lote = sofrito.lotes[0]
comprobar(lote.procedencia == "elaboración" and lote.fecha_caducidad == HOY + timedelta(days=4)
          and abs(lote.precio_unitario - (3 * 1.5 + 0.2 * 4) / 1.6) < 1e-3 and prep.coste == 5.3,
          "El lote del sofrito tiene su caducidad (vida útil 4 días) y su coste real (5,30 € / 1,6 kg)")
comprobar(Metricas(inv).gasto_por_categoria(HOY, HOY).get("Elaboraciones", 0) == 0,
          "Preparar una base no cuenta como compra (no es dinero gastado)")
silencio(inv.editar_producto, "Cebolla", "Cebolla blanca")
comprobar(sofrito.formula["ingredientes"] == {"Cebolla blanca": 1.5, "Aceite": 0.1},
          "Al renombrar un ingrediente, la fórmula de la base se actualiza")
copia = Inventario.from_dict(inv.to_dict())
comprobar(copia.buscar_producto("Sofrito").es_base() and copia.buscar_producto("Sofrito").vida_util_dias == 4
          and len(copia.elaboraciones.preparaciones_base) == 1
          and copia.elaboraciones.preparaciones_base[0].obtenida == 1.6,
          "Las bases, sus fórmulas y sus preparaciones se guardan y se cargan")

# Lista de la compra: un plato con salsa (que lleva sofrito, que lleva cebolla)
rec = Recetario()
guiso = Receta("Guiso", "Principales", {"Salsa": 0.1})
silencio(rec.agregar_receta, guiso)
silencio(rec.agregar_menu, Menu("Menú guiso", [guiso]))
serv = Servicio(HOY + timedelta(days=2), time(14, 0), 40, "Menú guiso")
compras = GestorCompras()
avisos = silencio(compras.generar_lista_desde_servicios, [serv], rec, inv)
pedidos = {i.ingrediente: i.cantidad for i in compras.items_pendientes()}
# 4 l de salsa -> 1 kg de sofrito (hay 1,6: no falta) + 2 kg de huesos (hay 3: no falta)
comprobar(pedidos == {} and any("Salsa" in a for a in avisos),
          "Si falta una base, la lista pide sus ingredientes (aquí hay de todo: no se compra nada)")
serv.comensales = 200  # 20 l de salsa -> 5 kg de sofrito (faltan 3,4) y 10 kg de huesos
silencio(compras.generar_lista_desde_servicios, [serv], rec, inv)
pedidos = {i.ingrediente: i.cantidad for i in compras.items_pendientes()}
nombre_cebolla = "Cebolla blanca"
comprobar("Salsa" not in pedidos and "Sofrito" not in pedidos and pedidos.get("Huesos") == 7
          and abs(pedidos.get(nombre_cebolla, 0) - (3.4 * 1.5 - 1)) < 1e-6,
          "Bases dentro de bases: no se compran salsa ni sofrito, sino huesos (7 kg) y la cebolla que falta (4,1 kg)")

with tempfile.TemporaryDirectory() as carpeta:
    from openpyxl import load_workbook
    hoja = load_workbook(silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta))["Elaboraciones base"]
    valores = [c.value for fila in hoja.iter_rows() for c in fila]
    comprobar("Sofrito" in valores and "Salsa" in valores and 1.6 in valores and "Obtenida" in valores,
              "El Excel tiene la hoja 'Elaboraciones base' con fórmulas y preparaciones")

print("\n--- Anotaciones (bloc de notas) ---")
from recetario import Menu as _Menu  # noqa: E402
r = Receta("Croquetas", "Entrantes", {"Aceite": 0.01})
comprobar(r.notas == "" and r.notas_fecha is None, "Una receta nueva no tiene notas")
r.poner_nota("  Bechamel del día anterior.\nFreír a 180 grados.  ")
comprobar(r.notas == "Bechamel del día anterior.\nFreír a 180 grados." and r.notas_fecha == HOY,
          "Se guarda la nota (con sus saltos de línea) y la fecha en que se editó")
r.poner_nota("Bechamel del día anterior.\nFreír a 180 grados.", fecha=HOY + timedelta(days=3))
comprobar(r.notas_fecha == HOY, "Guardar la misma nota sin cambios no cambia la fecha")
m = _Menu("Menú croquetas", [r])
m.poner_nota("Para eventos de pie.")
sofrito = inv.buscar_producto("Sofrito")
sofrito.poner_nota("Pochar sin dorar.")
rec = Recetario()
silencio(rec.agregar_receta, r)
silencio(rec.agregar_menu, m)
copia_rec = Recetario.from_dict(rec.to_dict())
copia_inv = Inventario.from_dict(inv.to_dict())
comprobar(copia_rec.recetas["Croquetas"].notas == r.notas and copia_rec.recetas["Croquetas"].notas_fecha == HOY
          and copia_rec.menus["Menú croquetas"].notas == "Para eventos de pie."
          and copia_inv.buscar_producto("Sofrito").notas == "Pochar sin dorar.",
          "Las notas de recetas, menús y bases se guardan y se cargan")
foto = m.foto()
comprobar(foto["notas"] == "Para eventos de pie." and foto["recetas"][0]["notas"] == r.notas,
          "La copia del menú de un servicio completado guarda también sus notas")
datos_viejos = r.to_dict()
del datos_viejos["notas"], datos_viejos["notas_fecha"]
comprobar(Receta.from_dict(datos_viejos).notas == "", "Las sesiones de antes de las notas se cargan sin notas")
r.poner_nota("")
comprobar(r.notas == "" and r.notas_fecha is None, "Una nota vacía la borra")
r.poner_nota("Bechamel del día anterior.")
with tempfile.TemporaryDirectory() as carpeta:
    from openpyxl import load_workbook
    libro = load_workbook(silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta, None, rec))
    valores = [c.value for fila in libro["Recetario"].iter_rows() for c in fila]
    bases = [c.value for fila in libro["Elaboraciones base"].iter_rows() for c in fila]
    comprobar("Croquetas" in valores and "Bechamel del día anterior." in valores and "Para eventos de pie." in valores
              and "Pochar sin dorar." in bases,
              "El Excel tiene la hoja 'Recetario' con recetas, menús y notas, y las notas de las bases")

print("\n--- Limpieza y mantenimiento ---")
from gastos import RegistroGastos as _RG, resumen_servicio as _resumen  # noqa: E402
inv = Inventario()
silencio(inv.agregar_producto, Producto("Cebolla", "Verduras", 5, "kg", 1, "Huerta"))
silencio(inv.agregar_producto, Producto("Lejía", "Limpieza", 4, "litros", 1.5, "Droguería", stock_minimo=2,
                                        tipo="mantenimiento", fecha_caducidad=HOY + timedelta(days=90)))
silencio(inv.agregar_producto, Producto("Bayetas", "Limpieza", 6, "unidades", 0.5, "Droguería", stock_minimo=10,
                                        tipo="mantenimiento"))
lejia = inv.buscar_producto("Lejía")
comprobar(inv.mantenimiento() == [lejia, inv.buscar_producto("Bayetas")] and lejia not in inv.alimentos()
          and lejia not in inv.consumibles() and lejia.fecha_caducidad == HOY + timedelta(days=90),
          "Los productos de limpieza y mantenimiento van en su propia lista y pueden tener caducidad")
try:
    Producto("Estropajo", "Limpieza", 1, "kg", 1, "Droguería", tipo="mantenimiento", tiene_merma=True)
    comprobar(False, "No pueden tener merma")
except ValueError:
    comprobar(True, "No pueden tener merma")
try:
    silencio(inv.definir_base, "Caldo raro", "Elaboraciones", "litros", 1, {"Lejía": 0.1})
    comprobar(False, "No pueden ser ingrediente")
except ValueError:
    comprobar("Caldo raro" not in inv.productos, "No pueden ser ingrediente de una elaboración base")

servicio = Servicio(HOY, time(20, 0), 10, "Menú inexistente")
servicio.id = 777
silencio(inv.salida_stock, "Lejía", 1, "consumo", 1, servicio_id=777)
silencio(inv.salida_stock, "Lejía", 0.5, "consumo", 1)
servicio.estado = "completado"
r = _resumen(servicio, inv, Recetario(), _RG())
comprobar(r["mantenimiento"] == 1.5 and r["comida"] == 0 and r["coste_total"] == 1.5,
          "Lo que sale para un servicio cuenta en su rentabilidad (1 l de lejía = 1,50 €); lo general no")
silencio(inv.entrada_stock, "Lejía", 2, precio_unitario=1.5, proveedor="Droguería")
silencio(inv.salida_stock, "Lejía", 2, "consumo", 2)  # se gasta lo comprado (para no cambiar el resto)
comprobar(Metricas(inv).gasto_por_tipo(HOY, HOY) == {"alimento": 0, "consumible": 0, "mantenimiento": 3.0},
          "Métricas separa el gasto en limpieza y mantenimiento (compra de 2 l de lejía = 3 €)")
compras = GestorCompras()
avisos = silencio(compras.generar_lista_desde_servicios, [], Recetario(), inv)
pedidos = {i.ingrediente: i.cantidad for i in compras.items_pendientes()}
comprobar(pedidos == {"Bayetas": 4} and any("Bayetas" in a and "mínimo" in a for a in avisos),
          "La lista de la compra repone hasta el mínimo lo que está por debajo (4 bayetas), aunque no haya servicios")
silencio(inv.salida_stock, "Lejía", 1, "consumo", 1)  # quedan 1,5 l de lejía (mínimo 2)
silencio(compras.generar_lista_desde_servicios, [], Recetario(), inv)
pedidos = {i.ingrediente: i.cantidad for i in compras.items_pendientes()}
comprobar(pedidos.get("Lejía") == 0.5 and pedidos.get("Bayetas") == 4,
          "Al volver a generarla no se borran: se mantienen y se añaden los nuevos (0,5 l de lejía)")
silencio(inv.editar_producto, "Bayetas", tipo="consumible")
comprobar(inv.buscar_producto("Bayetas").es_consumible() and inv.buscar_producto("Bayetas") in inv.consumibles(),
          "Se puede corregir el tipo si se dio de alta en la lista equivocada")
silencio(inv.editar_producto, "Bayetas", tipo="mantenimiento")
copia = Inventario.from_dict(inv.to_dict())
comprobar(copia.buscar_producto("Lejía").es_mantenimiento(), "Se guardan y se cargan")
with tempfile.TemporaryDirectory() as carpeta:
    from openpyxl import load_workbook
    hoja = load_workbook(silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta))["Inventario"]
    comprobar("Limpieza y mantenimiento" in [c.value for fila in hoja.iter_rows() for c in fila],
              "En el Excel aparecen con su clase")

print("\n--- Historial de precios ---")
from inventario import PrecioCompra  # noqa: E402
inv = Inventario()
silencio(inv.agregar_producto, Producto("Tomate", "Verduras", 5, "kg", 1.6, "Huerta"))
silencio(inv.entrada_stock, "Tomate", 4, precio_unitario=1.8, proveedor="Frutas Paco")
silencio(inv.entrada_stock, "Tomate", 3, precio_unitario=1.7, proveedor="Huerta")
precios = inv.precios_de("Tomate")
comprobar([(p.proveedor, p.precio_unitario, p.origen) for p in precios]
          == [("Huerta", 1.6, "inicial"), ("Frutas Paco", 1.8, "compra"), ("Huerta", 1.7, "compra")],
          "Cada compra (y el stock inicial) queda en el historial de precios, con su proveedor")
resumen = inv.resumen_precios_por_proveedor("Tomate")
comprobar(resumen[0]["proveedor"] == "Huerta" and abs(resumen[0]["medio"] - (5 * 1.6 + 3 * 1.7) / 8) < 1e-6
          and resumen[0]["compras"] == 2 and resumen[1]["proveedor"] == "Frutas Paco",
          "El resumen por proveedor da el precio medio ponderado, del más barato al más caro")
silencio(inv.salida_stock, "Tomate", 5, "consumo", 1)
comprobar(len(inv.precios_de("Tomate")) == 3, "Aunque el lote se gaste entero, su precio sigue en el historial")
comprobar(inv.avisos_precios() == [], "Sin cambios grandes de precio no hay avisos")
silencio(inv.entrada_stock, "Tomate", 2, precio_unitario=2.3, proveedor="Frutas Paco")
avisos = inv.avisos_precios()
comprobar(len(avisos) == 1 and avisos[0]["producto"] == "Tomate" and avisos[0]["variacion"] > 0.3,
          "Una compra mucho más cara de lo habitual da un aviso")
silencio(inv.entrada_stock, "Tomate", 2, precio_unitario=1.2, proveedor="Huerta")
avisos = inv.avisos_precios()
comprobar(len(avisos) == 1 and avisos[0]["variacion"] < -0.2 and avisos[0]["precio"] == 1.2,
          "...y una mucho más barata, también (solo cuenta la última compra de cada producto)")
comprobar(inv.avisos_precios(hoy=HOY + timedelta(days=40)) == [], "Las compras de hace más de 30 días ya no avisan")
silencio(inv.agregar_producto, Producto("Carne limpia", "Carnes", 0, "kg", 0, "Propio"))
limpio_antes = len(inv.historial_precios)
silencio(inv.entrada_stock, "Carne limpia", 1, precio_unitario=9, motivo="limpieza")
comprobar(len(inv.historial_precios) == limpio_antes, "Lo que sale de una limpieza o elaboración no es una compra: no entra")
silencio(inv.editar_producto, "Tomate", "Tomate pera")
copia = Inventario.from_dict(inv.to_dict())
comprobar(len(copia.precios_de("Tomate pera")) == 5 and not copia.precios_de("Tomate"),
          "Se renombra con el producto y se guarda y se carga")
viejo = inv.to_dict()
del viejo["historial_precios"]
reconstruido = Inventario.from_dict(viejo)
comprobar([(p.proveedor, p.precio_unitario) for p in reconstruido.precios_de("Tomate pera")]
          == [("Frutas Paco", 1.8), ("Huerta", 1.7), ("Frutas Paco", 2.3), ("Huerta", 1.2)],
          "En sesiones antiguas se reconstruye con las compras del historial de movimientos")
with tempfile.TemporaryDirectory() as carpeta:
    from openpyxl import load_workbook
    hoja = load_workbook(silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta))["Historial de precios"]
    comprobar(hoja.max_row == 6 and hoja["C2"].value in ("Huerta", "Frutas Paco"),
              "El Excel tiene la hoja 'Historial de precios' (una fila por compra)")

print("\n--- Corregir un precio hacia atrás ---")
from gastos import RegistroGastos as _RG2, resumen_servicio as _res2  # noqa: E402
inv = Inventario()
silencio(inv.agregar_producto, Producto("Lomo", "Carnes", 0, "kg", 0, "Carnicería"))
lote = silencio(inv.entrada_stock, "Lomo", 5, precio_unitario=14, proveedor="Carnicería")  # error: eran 4 €/kg
silencio(inv.salida_stock, "Lomo", 2, "consumo", lote.id, servicio_id=900)
silencio(inv.salida_stock, "Lomo", 0.5, "desperdicio", lote.id)
silencio(inv.salida_stock, "Lomo", 1, "elaboración", lote.id)
usos = inv.usos_del_lote("Lomo", lote.id)
comprobar(usos["compra"] is not None and usos["servicios"] == {900: 28.0} and len(usos["salidas"]) == 3
          and len(usos["derivados"]) == 1,
          "Se sabe qué hay registrado con un lote: su compra, el servicio #900, un desperdicio y una elaboración")
silencio(inv.editar_lote, "Lomo", lote.id, precio_unitario=4)
comprobar(lote.precio_unitario == 4 and usos["compra"].precio_unitario == 14 and inv.precios_de("Lomo")[0].precio_unitario == 14,
          "Sin corregir lo registrado, solo cambia el lote (lo que salga a partir de ahora)")
silencio(inv.editar_lote, "Lomo", lote.id, precio_unitario=4, proveedor="Carnes Paco", corregir_registrado=True)
servicio = Servicio(HOY, time(14, 0), 10, "Menú inexistente")
servicio.id = 900
servicio.estado = "completado"
comprobar(_res2(servicio, inv, Recetario(), _RG2())["comida"] == 8.0
          and Metricas(inv).gasto_por_categoria(HOY, HOY) == {"Carnes": 20.0}
          and inv.precios_de("Lomo")[0].precio_unitario == 4 and inv.precios_de("Lomo")[0].proveedor == "Carnes Paco"
          and "Carnes Paco" in usos["compra"].lote,
          "Corrigiendo lo registrado: el servicio cuesta 8 €, la compra 20 € y el historial dice 4 € (y el proveedor nuevo)")
silencio(inv.salida_stock, "Lomo", 1.5, "consumo", lote.id, servicio_id=901)
comprobar(inv.buscar_producto("Lomo").buscar_lote(lote.id) is None, "(el lote ya se ha gastado entero)")
compra = inv.precios_de("Lomo")[0]
silencio(inv.corregir_compra, compra, 4.5)
comprobar(compra.precio_unitario == 4.5 and inv.usos_del_lote("Lomo", lote.id)["servicios"] == {900: 9.0, 901: 6.75},
          "Desde el historial de precios se corrige una compra aunque su lote ya no exista")
try:
    inv.corregir_compra(compra, -1)
    comprobar(False, "Un precio negativo no se acepta")
except ValueError:
    comprobar(compra.precio_unitario == 4.5, "Un precio negativo no se acepta")

print("\n--- IVA ---")
from inventario import precio_a_coste  # noqa: E402
inv = Inventario()
silencio(inv.agregar_producto, Producto("Leche", "Lácteos", 0, "litros", 0, "Lácteos SA", iva=4))
comprobar(Producto("Detergente", "Limpieza", 0, "litros", 0, "Droguería").iva == 21, "El IVA por defecto es el 21 %")
comprobar(precio_a_coste(1.0, False, 4) == 1.04 and precio_a_coste(1.04, True, 4) == 1.04,
          "Un precio escrito sin IVA se convierte en lo pagado (1 € + 4 % = 1,04 €)")
lote = silencio(inv.entrada_stock, "Leche", 10, precio_unitario=1.04, proveedor="Lácteos SA")
comprobar(abs(lote.precio_base - 1.0) < 1e-9 and lote.iva == 4 and lote.precio_unitario == 1.04,
          "Se ve lo pagado (1,04 €/l) y se guarda aparte sin IVA (1 €) y su tipo (4 %)")
silencio(inv.salida_stock, "Leche", 3, "consumo", lote.id, servicio_id=950)
comprobar(Metricas(inv).iva_soportado(HOY, HOY) == 0.4 and Metricas(inv).gasto_por_categoria(HOY, HOY) == {"Lácteos": 10.4},
          "Métricas: el gasto es lo pagado (10,40 €) y el IVA soportado se ve aparte (0,40 €)")
servicio = Servicio(HOY, time(14, 0), 10, "Menú inexistente")
servicio.id = 950
servicio.estado = "completado"
servicio.precio_cobrado = 100
inv.iva_recuperable = False
r = _res2(servicio, inv, Recetario(), _RG2())
comprobar(r["comida"] == 3.12 and r["iva_recuperable"] == 0, "Si NO se recupera el IVA, el servicio cuesta lo pagado (3,12 €)")
inv.iva_recuperable = True
r = _res2(servicio, inv, Recetario(), _RG2())
comprobar(r["comida"] == 3.0 and r["iva_recuperable"] == 0.12 and r["margen"] == 97.0 and lote.precio_unitario == 1.04,
          "Si se recupera, el margen se calcula sin IVA (3 €) y el IVA recuperable va aparte (0,12 €); el precio se sigue viendo con IVA")
copia = Inventario.from_dict(inv.to_dict())
comprobar(copia.iva_recuperable and copia.buscar_producto("Leche").lotes[0].precio_unitario == 1.04
          and copia.historial[-1].iva == 4, "El ajuste y los precios se guardan y se cargan")
viejo = {"productos": [{"nombre": "Aceite", "categoria": "Despensa", "unidad": "litros", "proveedor": "X",
                        "stock_minimo": 0, "precio_referencia": 5, "lotes": [
                            {"id": 1, "cantidad": 2, "precio_unitario": 5, "proveedor": "X",
                             "fecha_entrada": HOY.isoformat(), "procedencia": "compra"}]}]}
antiguo = Inventario.from_dict(viejo)
comprobar(antiguo.buscar_producto("Aceite").precio_unitario == 5 and antiguo.buscar_producto("Aceite").lotes[0].iva == 0,
          "Las sesiones de antes del IVA se cargan con sus precios tal cual")

# Elaboraciones: conservan la parte de IVA de sus ingredientes
inv2 = Inventario()
silencio(inv2.agregar_producto, Producto("Cebolla", "Verduras", 3, "kg", 1.04, "Huerta", iva=4))
silencio(inv2.definir_base, "Sofrito", "Elaboraciones", "kg", 1, {"Cebolla": 1.5})
silencio(Recetario().preparar_base, "Sofrito", 1, 1, inv2)
sofrito = inv2.buscar_producto("Sofrito").lotes[0]
comprobar(abs(sofrito.precio_unitario - 1.56) < 1e-6 and abs(sofrito.precio_base - 1.5) < 1e-6,
          "Una elaboración base cuesta lo pagado (1,56 €) y guarda su parte sin IVA (1,50 €)")
rec2 = Recetario()
ensalada2 = Receta("Ensalada", "Entrantes", {"Cebolla": 0.1})
silencio(rec2.agregar_receta, ensalada2)
tanda = silencio(rec2.preparar_elaboracion, "Ensalada", 10, inv2)
comprobar(abs(tanda.coste_por_racion - 0.104) < 1e-6 and abs(tanda.coste_por_racion_sin_iva - 0.1) < 1e-6,
          "Una tanda guarda su coste con IVA y sin IVA")
silencio(inv2.editar_producto, "Cebolla", iva=10)
comprobar(inv2.buscar_producto("Cebolla").iva == 10 and inv2.buscar_producto("Cebolla").lotes[0].iva == 4,
          "Cambiar el IVA de un producto vale para las compras nuevas; las ya hechas conservan el suyo")

# IVA del trimestre (estimación)
from metricas import trimestre  # noqa: E402
registro = RegistroServicios()
servicio.fecha = HOY
silencio(registro.agregar_servicio, servicio)
desde, hasta = trimestre(HOY)
resumen = Metricas(inv).resumen_iva(desde, hasta, registro.servicios)
comprobar(resumen["soportado_por_tipo"] == {4.0: 0.4} and resumen["soportado"] == 0.4
          and resumen["repercutido"] == 10.0 and resumen["resultado"] == 9.6,
          "IVA del trimestre: 0,40 € pagado, 10 € cobrado (10 % de 100 €) -> 9,60 € a ingresar (estimación)")
with tempfile.TemporaryDirectory() as carpeta:
    from openpyxl import load_workbook
    libro = load_workbook(silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta))
    cabecera = [c.value for c in libro["Historial de precios"][1]]
    comprobar("IVA (%)" in cabecera and "Total con IVA (€)" in cabecera and libro["Historial de precios"]["G2"].value == 4,
              "El Excel separa precio sin IVA, IVA y precio con IVA")
AJUSTES["iva_recuperable"] = False

print("\n--- Fase 1: guardado seguro ---")
import json as _json  # noqa: E402
import os as _os  # noqa: E402
from persistencia import (  # noqa: E402
    guardar_sesion, cargar_sesion_segura, carpeta_datos, escribir_sesion, sesion_a_dict, firma,
)
from metricas import ArchivoInformes  # noqa: E402
with tempfile.TemporaryDirectory() as carpeta:
    ruta = Path(carpeta) / "sesion.json"
    inv = Inventario()
    silencio(inv.agregar_producto, Producto("Arroz", "Despensa", 5, "kg", 1.2, "Mayorista"))
    args = (inv, RegistroServicios(), Recetario(), GestorCompras(), ArchivoInformes())
    silencio(guardar_sesion, *args, str(ruta))
    silencio(inv.agregar_producto, Producto("Sal", "Despensa", 1, "kg", 0.5, "Mayorista"))
    silencio(guardar_sesion, *args, str(ruta))
    bak = ruta.with_name("sesion.json.bak")
    comprobar(ruta.exists() and bak.exists() and "Sal" not in bak.read_text(encoding="utf-8")
              and any((Path(carpeta) / "copias").glob("sesion_*.json")) and not ruta.with_name("sesion.json.tmp").exists(),
              "Guardar deja la versión anterior en .bak y una copia del día en copias/")
    ruta.write_text('{"inventario": {"productos": [', encoding="utf-8")  # archivo cortado a mitad
    sesion, aviso = silencio(cargar_sesion_segura, str(ruta))
    comprobar(sesion is not None and "Arroz" in sesion.inventario.productos and aviso and "recuperado" in aviso
              and any(Path(carpeta).glob("sesion_danada_*.json")),
              "Si sesion.json está dañado, se recupera de la copia y el dañado se aparta (no se borra)")
    for f in list(Path(carpeta).rglob("*.json*")):
        f.write_text("basura", encoding="utf-8")
    sesion, aviso = silencio(cargar_sesion_segura, str(ruta))
    comprobar(sesion is None and aviso and "vacío" in aviso, "Sin ninguna copia válida, se avisa y se empieza vacío")
    datos = sesion_a_dict(*args)
    ruta.write_text("\ufeff" + _json.dumps(datos), encoding="utf-8")  # con BOM (Bloc de notas)
    sesion, aviso = silencio(cargar_sesion_segura, str(ruta))
    comprobar(sesion is not None and aviso is None, "Un archivo guardado con BOM también se carga")
    comprobar(firma(datos) == firma(sesion_a_dict(*args)) and firma(datos) != firma({}),
              "La huella de los datos sirve para saber si hay cambios sin guardar")
from persistencia import SesionIlegible  # noqa: E402
import time as _time  # noqa: E402
with tempfile.TemporaryDirectory() as carpeta:
    ruta = Path(carpeta) / "sesion.json"
    inv = Inventario()
    args = (inv, RegistroServicios(), Recetario(), GestorCompras(), ArchivoInformes())
    silencio(inv.agregar_producto, Producto("Arroz", "Despensa", 5, "kg", 1.2, "Mayorista"))
    silencio(guardar_sesion, *args, str(ruta))
    _time.sleep(0.05)
    silencio(inv.agregar_producto, Producto("Sal", "Despensa", 1, "kg", 0.5, "Mayorista"))
    silencio(guardar_sesion, *args, str(ruta))  # .bak = solo arroz; copia de hoy = arroz + sal
    ruta.write_text("{roto", encoding="utf-8")
    sesion, aviso = silencio(cargar_sesion_segura, str(ruta))
    comprobar(sesion is not None and "Sal" in sesion.inventario.productos,
              "Al recuperar, se usa la copia MÁS RECIENTE (no se pierde lo último guardado)")
    datos = _json.loads(ruta.read_text(encoding="utf-8"))
    datos["inventario"]["productos"][0]["unidad"] = "cajas raras"  # JSON bueno, contenido que no se entiende
    ruta.write_text(_json.dumps(datos), encoding="utf-8")
    antes = ruta.read_text(encoding="utf-8")
    try:
        silencio(cargar_sesion_segura, str(ruta))
        comprobar(False, "Un contenido que no se entiende no se trata como archivo roto")
    except SesionIlegible:
        comprobar(ruta.read_text(encoding="utf-8") == antes,
                  "Si el archivo se lee pero no se entiende, se avisa y NO se aparta ni se sobrescribe")
with tempfile.TemporaryDirectory() as carpeta:
    ruta = Path(carpeta) / "sesion.json"
    (Path(carpeta) / "copias").write_text("no soy una carpeta", encoding="utf-8")  # las copias no se pueden hacer
    silencio(escribir_sesion, {"a": 1}, str(ruta))
    comprobar(_json.loads(ruta.read_text(encoding="utf-8")) == {"a": 1},
              "Si la copia del día falla, el guardado principal se hace igual")
a1, a2 = Inventario(), Inventario()
a1.iva_cobro = 21
comprobar(a2.iva_cobro == 10, "Los ajustes de IVA son de cada sesión (dos ventanas no se los pisan)")
_os.environ["GESTION_RESTAURANTE_DATOS"] = "/tmp/otra"
comprobar(carpeta_datos(Path("/x")) == Path("/tmp/otra"), "La carpeta de datos se puede fijar (pruebas)")
del _os.environ["GESTION_RESTAURANTE_DATOS"]
comprobar(carpeta_datos(Path("/x")) == Path("/x/datos"), "Fuera del .exe, los datos siguen en la carpeta 'datos'")

print("\n--- Fase 1: servicios, recetas y limpiezas ---")
registro = RegistroServicios()
s1 = Servicio(HOY, time(14, 0), 10, "Menú A")
s2 = Servicio(HOY, time(14, 0), 10, "Menú A")
s2.id = s1.id  # como si otra ventana hubiera dado el mismo número
silencio(registro.agregar_servicio, s1)
silencio(registro.agregar_servicio, s2)
comprobar(s1.id != s2.id, "Dos servicios nunca tienen el mismo número")
s1.completar()
try:
    registro.cancelar_servicio(s1.id)
    comprobar(False, "Un servicio completado no se puede cancelar")
except ValueError:
    comprobar(s1.estado == "completado", "Un servicio completado no se puede cancelar")
rec = Recetario()
silencio(rec.agregar_receta, Receta("Paella", "Arroces", {"Arroz": 0.1}))
for nombre in ("Paella", "paella ", " PAELLA"):
    try:
        rec.agregar_receta(Receta(nombre, "Arroces", {"Arroz": 0.2}))
        comprobar(False, f"No se puede crear otra receta '{nombre}'")
    except ValueError:
        pass
comprobar(rec.recetas["Paella"].ingredientes_por_comensal == {"Arroz": 0.1},
          "Una receta con un nombre que ya existe (aunque cambien mayúsculas o espacios) no sustituye a la otra")
silencio(rec.agregar_menu, Menu("Menú A", [rec.recetas["Paella"]]))
try:
    rec.agregar_menu(Menu("menú a", []))
    comprobar(False, "No se puede repetir un menú")
except ValueError:
    comprobar(True, "Un menú con un nombre que ya existe no sustituye al otro")
avisos = silencio(GestorCompras().generar_lista_desde_servicios,
                  [Servicio(HOY + timedelta(days=1), time(14, 0), 10, "Menu A")], rec, Inventario())
comprobar(any("no existe" in a for a in avisos), "La lista de la compra AVISA de un servicio cuyo menú no existe")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Pata", "Carnes", 2, "unidades", 40, "Carnicería", tiene_merma=True, peso_unitario=7))
silencio(inv.agregar_producto, Producto("Cebolla", "Verduras", 5, "kg", 1, "Huerta"))
silencio(inv.definir_base, "Fondo", "Elaboraciones", "litros", 1, {"Cebolla": 1})
silencio(inv.definir_base, "Sofrito", "Elaboraciones", "kg", 1, {"Cebolla": 1})
try:
    silencio(inv.limpiar_producto, "Pata", 1, "Sofrito", 5, lote_id=1)
    comprobar(False, "Una limpieza no puede dar una elaboración base")
except ValueError:
    copia = Inventario.from_dict(inv.to_dict())
    comprobar(inv.buscar_producto("Pata").stock == 2 and copia.buscar_producto("Sofrito").es_base(),
              "Una limpieza no puede dar una elaboración base (antes dejaba la sesión imposible de abrir)")

print()
if fallos:
    print(f"RESULTADO: {len(fallos)} FALLO(S)")
    sys.exit(1)
print("RESULTADO: TODO BIEN")
