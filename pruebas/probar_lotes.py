"""
pruebas/probar_lotes.py
-------------------------
Prueba la LÓGICA de los lotes (sin interfaz): entradas, salidas de un lote
elegido, desechar lotes caducados, limpiezas por lote, completar servicios
eligiendo lotes, guardar/cargar y la conversión de sesiones antiguas.

Es un script normal de Python (sin pytest). Se ejecuta en GitHub Actions
junto a las pruebas de la interfaz, y también a mano:
    python pruebas/probar_lotes.py
"""

import io
import sys
import tempfile
from contextlib import redirect_stdout
from datetime import date, time, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "modulos"))

from inventario import Inventario, Producto  # noqa: E402
from metricas import Metricas  # noqa: E402
from recetario import Menu, Receta, Recetario  # noqa: E402
from servicios import Servicio  # noqa: E402
from exportador import exportar_todo  # noqa: E402
from servicios import RegistroServicios  # noqa: E402
from compras import GestorCompras  # noqa: E402

HOY = date.today()
fallos: list[str] = []


def comprobar(condicion: bool, descripcion: str) -> None:
    print(("✅ " if condicion else "❌ ") + descripcion)
    if not condicion:
        fallos.append(descripcion)


def silencio(funcion, *args, **kwargs):
    """Ejecuta sin llenar la pantalla con los print() de los módulos."""
    with redirect_stdout(io.StringIO()):
        return funcion(*args, **kwargs)


def inventario_secreto() -> Inventario:
    """Tu ejemplo: 1 kg de secreto que caduca pronto + una compra nueva que caduca más tarde."""
    inv = Inventario()
    silencio(inv.agregar_producto, Producto(
        "Secreto", "Carnes", 1, "kg", 14, "Carnicería Pepe", stock_minimo=0.5,
        fecha_caducidad=HOY + timedelta(days=10),
    ))
    silencio(inv.entrada_stock, "Secreto", 0.8, precio_unitario=15, proveedor="Ibéricos Sierra",
             fecha_caducidad=HOY + timedelta(days=15))
    return inv


print("--- Entradas: cada compra es un lote ---")
inv = inventario_secreto()
secreto = inv.buscar_producto("Secreto")
comprobar(len(secreto.lotes) == 2, "Una compra nueva crea un segundo lote")
comprobar(secreto.stock == 1.8, "El stock total es la suma de los lotes (1,8 kg)")
comprobar(secreto.fecha_caducidad == HOY + timedelta(days=10), "La caducidad del producto es la del lote más próximo")
comprobar(abs(secreto.valor_total() - 26) < 1e-9, "El valor es cada lote a su precio (14 + 12 = 26 €)")
comprobar(secreto.lotes[1].proveedor == "Ibéricos Sierra", "Cada lote guarda su proveedor")
silencio(inv.entrada_stock, "Secreto", 0.5, fecha_caducidad=HOY + timedelta(days=15))
comprobar(len(secreto.lotes) == 3, "Aunque tenga la misma caducidad, cada compra es un lote aparte")
comprobar(secreto.lotes[2].proveedor == "Carnicería Pepe" and secreto.lotes[2].precio_unitario == 15,
          "Sin indicar proveedor ni precio se usan el habitual y el de la última compra")
silencio(inv.agregar_producto, Producto("Aceite", "Aceites", 5, "litros", 4, "Oleícola"))
silencio(inv.entrada_stock, "Aceite", 5, precio_unitario=4.5)
aceite = inv.buscar_producto("Aceite")
comprobar(len(aceite.lotes) == 2 and aceite.fecha_caducidad is None,
          "Los productos sin caducidad también tienen un lote por compra")

print("\n--- Salidas: del lote que elige el usuario ---")
inv = inventario_secreto()
secreto = inv.buscar_producto("Secreto")
comprobar(silencio(inv.salida_stock, "Secreto", 0.3, "consumo", 2), "Se puede sacar del lote que caduca MÁS TARDE")
comprobar(abs(secreto.buscar_lote(2).cantidad - 0.5) < 1e-9 and secreto.buscar_lote(1).cantidad == 1,
          "Solo baja el lote elegido")
comprobar(not silencio(inv.salida_stock, "Secreto", 0.6, "consumo", 2), "No se puede sacar de un lote más de lo que tiene")
comprobar(inv.historial[-1].lote_id == 2 and inv.historial[-1].precio_unitario == 15,
          "El historial apunta el lote y su precio real")
silencio(inv.salida_stock, "Secreto", 0.5, "consumo", 2)
comprobar(secreto.buscar_lote(2) is None and len(secreto.lotes) == 1, "Un lote que llega a 0 desaparece")
silencio(inv.entrada_stock, "Secreto", 1)
comprobar(secreto.lotes[-1].id == 3, "Los números de lote no se reutilizan")

print("\n--- Desechar un lote caducado ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Tomate", "Verduras", 2, "kg", 2, "Huerta", fecha_caducidad=HOY - timedelta(days=1)))
silencio(inv.entrada_stock, "Tomate", 3, precio_unitario=2.5, fecha_caducidad=HOY + timedelta(days=5))
comprobar(len(inv.lotes_caducados()) == 1, "Se detecta el lote caducado")
comprobar(len(inv.lotes_proximos_a_caducar()) == 1, "Y el que caduca en los próximos días, aparte")
comprobar(silencio(inv.desechar_lote, "Tomate", 1), "Se desecha el lote caducado")
tomate = inv.buscar_producto("Tomate")
comprobar(tomate.stock == 3 and not inv.lotes_caducados(), "Solo queda el lote bueno")
desperdicio = Metricas(inv).cantidad_desperdiciada("Tomate", HOY, HOY)
comprobar(desperdicio == 2, "Queda registrado como desperdicio (2 kg)")
comprobar(Metricas(inv).gasto_por_categoria(HOY, HOY) == {"Verduras": 7.5},
          "El gasto cuenta la compra a su precio (3 kg x 2,5 €); el stock inicial no es una compra")

print("\n--- Limpieza de un lote concreto ---")
inv = Inventario()
silencio(inv.agregar_producto, Producto("Pata", "Carnes", 2, "unidades", 45, "Carnicería Pepe",
                                        tiene_merma=True, peso_unitario=7))
silencio(inv.entrada_stock, "Pata", 1, precio_unitario=50, peso_unitario=8)
pata = inv.buscar_producto("Pata")
comprobar(pata.lotes[1].peso_unitario == 8, "Cada lote de patas guarda su propio peso por unidad")
try:
    silencio(inv.limpiar_producto, "Pata", 1, "Carne limpia", 5)
    comprobar(False, "Con varios lotes hay que elegir cuál se limpia")
except ValueError:
    comprobar(True, "Con varios lotes hay que elegir cuál se limpia")
limpieza = silencio(inv.limpiar_producto, "Pata", 1, "Carne limpia", 5, {"Hueso": 1.5}, lote_id=2)
comprobar(limpieza.peso_bruto_kg == 8 and limpieza.coste == 50, "Usa el peso y el precio de ESE lote (8 kg, 50 €)")
comprobar(limpieza.lote_origen and "Lote 2" in limpieza.lote_origen, "La limpieza recuerda de qué lote salió")
limpia = inv.buscar_producto("Carne limpia")
comprobar(limpia.lotes[0].precio_unitario == 10, "El lote limpio carga con todo el coste (50 € / 5 kg)")
comprobar(inv.buscar_producto("Hueso").lotes[0].precio_unitario == 0, "Los derivados entran a coste 0")
comprobar(pata.stock == 2 and pata.buscar_lote(2) is None, "Solo sale el lote limpiado")

print("\n--- Completar un servicio eligiendo lotes ---")
inv = inventario_secreto()
recetario = Recetario()
silencio(recetario.agregar_receta, Receta("Secreto a la brasa", "Principal", {"Secreto": 0.2}))
silencio(recetario.agregar_menu, Menu("Menú brasa", [recetario.recetas["Secreto a la brasa"]]))
servicio = Servicio(HOY, time(21, 0), 5, "Menú brasa")  # 1 kg de secreto

fila = recetario.previsualizar_consumo(servicio, inv)[0]
comprobar(fila["lotes_elegidos"] == [1] and fila["sin_asignar"] == 0,
          "Por defecto propone el lote que caduca antes (y aquí llega)")
fila = recetario.previsualizar_consumo(servicio, inv, {"Secreto": [2]})[0]
comprobar(abs(fila["sin_asignar"] - 0.2) < 1e-9 and fila["lotes_restantes"] == [1],
          "Si el lote elegido no llega, dice cuánto falta y con qué lotes se puede completar")
try:
    silencio(recetario.completar_servicio, servicio, inv, {"Secreto": [2]})
    comprobar(False, "No deja completar sin elegir con qué lote se completa")
except ValueError:
    comprobar(True, "No deja completar sin elegir con qué lote se completa")
comprobar(inv.buscar_producto("Secreto").stock == 1.8 and servicio.estado == "pendiente", "...y no toca nada")
silencio(recetario.completar_servicio, servicio, inv, {"Secreto": [2, 1]})
secreto = inv.buscar_producto("Secreto")
comprobar(secreto.buscar_lote(2) is None and abs(secreto.buscar_lote(1).cantidad - 0.8) < 1e-9,
          "Saca 0,8 del lote 2 y los 0,2 que faltan del lote 1")
comprobar(servicio.estado == "completado", "El servicio queda completado")

inv = inventario_secreto()
grande = Servicio(HOY, time(21, 0), 20, "Menú brasa")  # 4 kg: no hay suficiente
fila = recetario.previsualizar_consumo(grande, inv, {"Secreto": [2]})[0]
comprobar(fila["sin_asignar"] == 0 and fila["faltante"] == 2.2, "Si no hay bastante en total, se usan todos los lotes")
silencio(recetario.completar_servicio, grande, inv, {"Secreto": [2]})
comprobar(inv.buscar_producto("Secreto").stock == 0, "...y se descuenta todo lo que había")

print("\n--- Guardar, cargar y sesiones antiguas ---")
inv = inventario_secreto()
silencio(inv.salida_stock, "Secreto", 0.3, "consumo", 1)
copia = Inventario.from_dict(inv.to_dict())
s2 = copia.buscar_producto("Secreto")
comprobar([(l.id, l.cantidad, l.proveedor) for l in s2.lotes] == [(1, 0.7, "Carnicería Pepe"), (2, 0.8, "Ibéricos Sierra")],
          "Los lotes se guardan y se cargan igual")
comprobar(copia.historial[-1].lote == inv.historial[-1].lote, "El historial conserva el lote de cada movimiento")
antiguo = {"productos": [{
    "nombre": "Harina", "categoria": "Panadería", "stock": 3, "unidad": "kg", "precio_unitario": 1.2,
    "proveedor": "Harinas del Sur", "stock_minimo": 2, "fecha_caducidad": (HOY + timedelta(days=30)).isoformat(),
}], "historial": [{"producto_nombre": "Harina", "categoria": "Panadería", "tipo": "entrada", "cantidad": 3,
                   "unidad": "kg", "precio_unitario": 1.2, "fecha": HOY.isoformat()}]}
viejo = Inventario.from_dict(antiguo)
harina = viejo.buscar_producto("Harina")
comprobar(len(harina.lotes) == 1 and harina.stock == 3 and harina.precio_unitario == 1.2
          and harina.fecha_caducidad == HOY + timedelta(days=30),
          "Una sesión de antes de los lotes se convierte en un lote con su stock, precio y caducidad")

print("\n--- Corregir lotes y productos ---")
inv = inventario_secreto()
silencio(inv.editar_lote, "Secreto", 2, cantidad=0.75, fecha_caducidad=HOY + timedelta(days=20), proveedor="Otro")
lote = inv.buscar_producto("Secreto").buscar_lote(2)
comprobar(lote.cantidad == 0.75 and lote.proveedor == "Otro" and lote.fecha_caducidad == HOY + timedelta(days=20),
          "Se puede corregir la cantidad, el proveedor y la caducidad de un lote")
silencio(inv.editar_lote, "Secreto", 2, borrar_fecha_caducidad=True)
comprobar(lote.fecha_caducidad is None, "Se puede quitar la caducidad de un lote")
silencio(inv.editar_producto, "Secreto", nuevo_nombre="Secreto ibérico", stock_minimo=1)
comprobar(inv.buscar_producto("Secreto ibérico").stock == 1.75, "Renombrar el producto conserva sus lotes")

print("\n--- Excel ---")
inv = inventario_secreto()
with tempfile.TemporaryDirectory() as carpeta:
    ruta = silencio(exportar_todo, inv, RegistroServicios(), GestorCompras(), carpeta)
    from openpyxl import load_workbook
    libro = load_workbook(ruta)
    comprobar("Lotes" in libro.sheetnames and libro["Lotes"].max_row == 4,
              "El Excel tiene una hoja 'Lotes' con un renglón por lote (+ cabecera y total)")

print()
if fallos:
    print(f"RESULTADO: {len(fallos)} FALLO(S)")
    sys.exit(1)
print("RESULTADO: TODO BIEN")
